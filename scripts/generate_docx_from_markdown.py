from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_table(document: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if cells and all(set(cell) <= {"-", ":"} for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        table.rows[0].cells[index].text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row[: len(cells)]):
            cells[index].text = value


def add_paragraph(document: Document, line: str, in_code: bool = False) -> None:
    if in_code:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        return

    stripped = line.strip()
    if not stripped:
        document.add_paragraph()
        return
    if stripped.startswith("- "):
        document.add_paragraph(stripped[2:], style="List Bullet")
        return
    if stripped[0:2].isdigit() and ". " in stripped[:5]:
        document.add_paragraph(stripped.split(". ", 1)[1], style="List Number")
        return
    document.add_paragraph(stripped.replace("`", ""))


def markdown_to_docx(source: Path, destination: Path) -> None:
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    lines = source.read_text(encoding="utf-8").splitlines()
    table_lines: list[str] = []
    in_code = False

    def flush_table() -> None:
        nonlocal table_lines
        if table_lines:
            add_table(document, table_lines)
            table_lines = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_table()
            in_code = not in_code
            continue
        if in_code:
            add_paragraph(document, line, in_code=True)
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines.append(line)
            continue
        flush_table()
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped == "---":
            document.add_page_break()
        else:
            add_paragraph(document, line)

    flush_table()
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a DOCX copy from a Markdown document.")
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    markdown_to_docx(args.source, args.destination)


if __name__ == "__main__":
    main()
