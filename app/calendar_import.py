from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app import models


MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}


@dataclass(frozen=True)
class CalendarImportRow:
    sheet: str
    row_number: int
    year: int
    date_label: str
    event_name: str
    start_date: date
    end_date: date


def normalize_calendar_event_name(value: str) -> str:
    normalized = value.strip().lower()
    normalized = normalized.replace("&", "and")
    normalized = normalized.replace("’", "'")
    normalized = re.sub(r"\s+", " ", normalized)
    normalized = re.sub(r"\s+([()])", r"\1", normalized)
    normalized = re.sub(r"([()])\s+", r"\1", normalized)
    return normalized


def calendar_event_name_candidates(value: str, year: int) -> set[str]:
    normalized = normalize_calendar_event_name(value)
    candidates = _event_name_semantic_variants(normalized)
    year_patterns = [
        rf"\s*\({year}\)\s*$",
        rf"\s*\({year}\s+season\)\s*$",
        rf"\s*{year}\s*$",
    ]
    for pattern in year_patterns:
        for candidate in list(candidates):
            stripped = re.sub(pattern, "", candidate).strip()
            if stripped:
                candidates.update(_event_name_semantic_variants(stripped))
    return candidates


def _event_name_semantic_variants(normalized: str) -> set[str]:
    variants = {normalized}
    discipline_suffix_patterns = [
        r"\s*\((?:mag|wag|mag and wag)\)\s*$",
        r"\s+(?:mag|wag)\s*$",
    ]
    for pattern in discipline_suffix_patterns:
        stripped = re.sub(pattern, "", normalized).strip()
        if stripped:
            variants.add(stripped)

    for candidate in list(variants):
        semantic = _replace_discipline_words(candidate)
        if semantic:
            variants.add(semantic)
        compact_parenthetical = re.sub(r"\((mag|wag)\)", r"\1", candidate).strip()
        if compact_parenthetical:
            variants.add(compact_parenthetical)
            semantic_compact = _replace_discipline_words(compact_parenthetical)
            if semantic_compact:
                variants.add(semantic_compact)

    return {candidate for candidate in variants if candidate}


def _replace_discipline_words(value: str) -> str:
    normalized = re.sub(r"\bmen's\b|\bmens\b|\bmen\b", "mag", value)
    normalized = re.sub(r"\bwomen's\b|\bwomens\b|\bwomen\b", "wag", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def parse_calendar_date_label(label: str, year: int) -> tuple[date, date]:
    raw = str(label).strip()
    if not raw:
        raise ValueError("DATE is required")
    normalized = raw.replace("–", "-").replace("—", "-").replace("−", "-")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    normalized = re.sub(r"\s*-\s*", "-", normalized)
    normalized = normalized.replace(",", "")

    month_pattern = r"([A-Za-z.]+)"
    day_pattern = r"(\d{1,2})"

    cross_month = re.fullmatch(
        rf"{month_pattern}\s+{day_pattern}-{month_pattern}\s+{day_pattern}",
        normalized,
    )
    if cross_month:
        start_month = _parse_month(cross_month.group(1))
        start_day = int(cross_month.group(2))
        end_month = _parse_month(cross_month.group(3))
        end_day = int(cross_month.group(4))
        end_year = year + 1 if end_month < start_month else year
        return date(year, start_month, start_day), date(end_year, end_month, end_day)

    same_month_range = re.fullmatch(
        rf"{month_pattern}\s+{day_pattern}-{day_pattern}",
        normalized,
    )
    if same_month_range:
        month = _parse_month(same_month_range.group(1))
        start_day = int(same_month_range.group(2))
        end_day = int(same_month_range.group(3))
        return date(year, month, start_day), date(year, month, end_day)

    single_day = re.fullmatch(rf"{month_pattern}\s+{day_pattern}", normalized)
    if single_day:
        month = _parse_month(single_day.group(1))
        day = int(single_day.group(2))
        single_date = date(year, month, day)
        return single_date, single_date

    raise ValueError(f"Unsupported DATE format: {label}")


def _parse_month(value: str) -> int:
    key = value.strip().strip(".").lower()
    month = MONTHS.get(key)
    if month is None:
        raise ValueError(f"Unsupported month: {value}")
    return month


def parse_calendar_file(filename: str, content: bytes) -> tuple[list[CalendarImportRow], list[dict]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _parse_calendar_csv(content)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_calendar_xlsx(content)
    raise ValueError("Calendar import supports .xlsx, .xlsm or .csv files")


def _parse_calendar_xlsx(content: bytes) -> tuple[list[CalendarImportRow], list[dict]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows: list[CalendarImportRow] = []
    issues: list[dict] = []
    for worksheet in workbook.worksheets:
        try:
            year = int(str(worksheet.title).strip())
        except ValueError:
            issues.append({
                "severity": "warning",
                "sheet": worksheet.title,
                "message": "Sheet skipped because its name is not a year",
            })
            continue

        header_date = str(worksheet.cell(row=1, column=1).value or "").strip().upper()
        header_event = str(worksheet.cell(row=1, column=2).value or "").strip().upper()
        if header_date != "DATE" or header_event != "EVENT":
            issues.append({
                "severity": "error",
                "sheet": worksheet.title,
                "message": "Expected headers DATE and EVENT in columns A and B",
            })
            continue

        for row_number in range(2, worksheet.max_row + 1):
            date_label = worksheet.cell(row=row_number, column=1).value
            event_name = worksheet.cell(row=row_number, column=2).value
            if date_label is None and event_name is None:
                continue
            if not event_name:
                issues.append({
                    "severity": "error",
                    "sheet": worksheet.title,
                    "row": row_number,
                    "message": "EVENT is required",
                })
                continue
            try:
                start_date, end_date = parse_calendar_date_label(str(date_label), year)
            except Exception as exc:
                issues.append({
                    "severity": "error",
                    "sheet": worksheet.title,
                    "row": row_number,
                    "event_name": str(event_name).strip(),
                    "message": str(exc),
                })
                continue
            rows.append(CalendarImportRow(
                sheet=worksheet.title,
                row_number=row_number,
                year=year,
                date_label=str(date_label).strip(),
                event_name=str(event_name).strip(),
                start_date=start_date,
                end_date=end_date,
            ))
    workbook.close()
    return rows, issues


def _parse_calendar_csv(content: bytes) -> tuple[list[CalendarImportRow], list[dict]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    required = {"DATE", "EVENT", "YEAR"}
    headers = {header.strip().upper() for header in (reader.fieldnames or [])}
    if not required.issubset(headers):
        return [], [{
            "severity": "error",
            "message": "CSV calendar import requires DATE, EVENT and YEAR columns",
        }]

    rows: list[CalendarImportRow] = []
    issues: list[dict] = []
    for row_number, item in enumerate(reader, start=2):
        normalized_item = {key.strip().upper(): value for key, value in item.items() if key}
        try:
            year = int(str(normalized_item.get("YEAR", "")).strip())
            event_name = str(normalized_item.get("EVENT", "")).strip()
            date_label = str(normalized_item.get("DATE", "")).strip()
            if not event_name:
                raise ValueError("EVENT is required")
            start_date, end_date = parse_calendar_date_label(date_label, year)
        except Exception as exc:
            issues.append({
                "severity": "error",
                "row": row_number,
                "message": str(exc),
            })
            continue
        rows.append(CalendarImportRow(
            sheet=str(year),
            row_number=row_number,
            year=year,
            date_label=date_label,
            event_name=event_name,
            start_date=start_date,
            end_date=end_date,
        ))
    return rows, issues


def infer_event_discipline(event_name: str) -> models.EventDisciplineEnum:
    upper_name = event_name.upper()
    has_mag = "(MAG)" in upper_name
    has_wag = "(WAG)" in upper_name
    if has_mag and not has_wag:
        return models.EventDisciplineEnum.MAG
    if has_wag and not has_mag:
        return models.EventDisciplineEnum.WAG
    return models.EventDisciplineEnum.MAG_AND_WAG


def infer_event_category(event_name: str) -> models.EventCategoryEnum:
    lower_name = event_name.lower()
    has_junior = "junior" in lower_name
    has_senior = "senior" in lower_name
    if has_junior and not has_senior:
        return models.EventCategoryEnum.JUNIOR
    if has_senior and not has_junior:
        return models.EventCategoryEnum.SENIOR
    return models.EventCategoryEnum.JUNIOR_AND_SENIOR


def infer_event_level(event_name: str) -> models.LevelEnum:
    lower_name = event_name.lower()
    if "olympic" in lower_name:
        return models.LevelEnum.OLYMPIC_GAMES
    if "world challenge cup" in lower_name:
        return models.LevelEnum.WORLD_CHALLENGE_CUP
    if "world cup" in lower_name:
        return models.LevelEnum.WORLD_CUP
    if "world championships" in lower_name:
        return models.LevelEnum.WORLD_CHAMPIONSHIPS
    if any(token in lower_name for token in ("european", "asian", "african", "pan american", "continental")):
        return models.LevelEnum.CONTINENTAL_CHAMPIONSHIPS
    if any(token in lower_name for token in ("national", "championships", "championship")):
        return models.LevelEnum.NATIONAL_EVENT
    return models.LevelEnum.INTERNATIONAL_EVENT


def summarize_calendar_import(
    db: Session,
    rows: list[CalendarImportRow],
    issues: list[dict],
    create_missing_from_year: int,
) -> dict:
    event_lookup = _build_event_lookup(db, {row.year for row in rows})
    seen_source_keys: dict[tuple[int, str], CalendarImportRow] = {}

    preview_rows = []
    duplicate_source_rows = []
    matched_event_ids: set[int] = set()
    would_update_event_ids: set[int] = set()
    already_up_to_date_event_ids: set[int] = set()
    would_create_source_keys: set[tuple[int, str]] = set()
    unmatched_historical_rows = []
    matched_sources_by_event_id: dict[int, list[CalendarImportRow]] = {}

    for row in rows:
        source_key = (row.year, normalize_calendar_event_name(row.event_name))
        duplicate_of = seen_source_keys.get(source_key)
        if duplicate_of is not None:
            duplicate_source_rows.append(_build_duplicate_source_row(row, duplicate_of))
        else:
            seen_source_keys[source_key] = row

        matches = _find_existing_events(event_lookup, row)
        matched_ids = [event.id for event in matches]
        matched_event_ids.update(matched_ids)
        for event in matches:
            matched_sources_by_event_id.setdefault(event.id, []).append(row)
        events_to_update = [
            event for event in matches
            if event.start_date != row.start_date or event.end_date != row.end_date
        ]
        would_update_event_ids.update(event.id for event in events_to_update)
        already_up_to_date_event_ids.update(
            event.id for event in matches
            if event.id not in {item.id for item in events_to_update}
        )

        if matches:
            action = "update_dates" if events_to_update else "no_change"
            match_status = "matched_multiple" if len(matches) > 1 else "matched"
        elif row.year >= create_missing_from_year:
            action = "create_event"
            match_status = "new_future_event"
            would_create_source_keys.add(source_key)
        else:
            action = "skip_unmatched_historical"
            match_status = "unmatched_historical"
            unmatched_historical_rows.append(_row_preview(row, [], match_status, action))

        preview_rows.append(_row_preview(row, matches, match_status, action))

    return {
        "parsed_rows": len(rows),
        "years": sorted({row.year for row in rows}),
        "matched_rows": sum(1 for row in preview_rows if row["matched_event_ids"]),
        "matched_events": len(matched_event_ids),
        "would_update_events": len(would_update_event_ids),
        "already_up_to_date_events": len(already_up_to_date_event_ids - would_update_event_ids),
        "would_create_events": len(would_create_source_keys),
        "unmatched_historical_rows": len(unmatched_historical_rows),
        "duplicate_source_rows": duplicate_source_rows,
        "matched_event_source_conflicts": _build_matched_event_source_conflicts(matched_sources_by_event_id),
        "issues": issues,
        "sample_rows": preview_rows[:25],
        "rows": preview_rows,
    }


def _build_event_lookup(db: Session, years: set[int]) -> dict[tuple[int, str], list[models.Event]]:
    query = db.query(models.Event).filter(models.Event.is_deleted.is_(False))
    if years:
        query = query.filter(models.Event.year.in_(years))
    lookup: dict[tuple[int, str], list[models.Event]] = {}
    for event in query.all():
        for candidate in calendar_event_name_candidates(event.name, event.year):
            lookup.setdefault((event.year, candidate), []).append(event)
    return lookup


def _find_existing_events(
    event_lookup: dict[tuple[int, str], list[models.Event]],
    row: CalendarImportRow,
) -> list[models.Event]:
    matches_by_id = {}
    for candidate in calendar_event_name_candidates(row.event_name, row.year):
        for event in event_lookup.get((row.year, candidate), []):
            matches_by_id[event.id] = event
    return sorted(matches_by_id.values(), key=lambda event: event.id)


def _row_preview(
    row: CalendarImportRow,
    matches: list[models.Event],
    match_status: str,
    action: str,
) -> dict:
    return {
        "sheet": row.sheet,
        "row": row.row_number,
        "year": row.year,
        "date_label": row.date_label,
        "event_name": row.event_name,
        "start_date": row.start_date,
        "end_date": row.end_date,
        "matched_event_ids": [event.id for event in matches],
        "matched_event_names": [event.name for event in matches],
        "match_status": match_status,
        "action": action,
        "inferred_discipline": infer_event_discipline(row.event_name).value,
        "inferred_category": infer_event_category(row.event_name).value,
        "inferred_level": infer_event_level(row.event_name).value,
    }


def _build_duplicate_source_row(
    row: CalendarImportRow,
    duplicate_of: CalendarImportRow,
) -> dict:
    return {
        "sheet": row.sheet,
        "row": row.row_number,
        "event_name": row.event_name,
        "year": row.year,
        "date_label": row.date_label,
        "duplicate_of_sheet": duplicate_of.sheet,
        "duplicate_of_row": duplicate_of.row_number,
        "duplicate_of_date_label": duplicate_of.date_label,
    }


def _build_matched_event_source_conflicts(
    matched_sources_by_event_id: dict[int, list[CalendarImportRow]],
) -> list[dict]:
    conflicts = []
    for event_id, source_rows in matched_sources_by_event_id.items():
        if len(source_rows) < 2:
            continue
        date_ranges = {
            (row.start_date, row.end_date)
            for row in source_rows
        }
        if len(date_ranges) <= 1:
            continue
        conflicts.append({
            "event_id": event_id,
            "source_rows": [
                {
                    "sheet": row.sheet,
                    "row": row.row_number,
                    "year": row.year,
                    "event_name": row.event_name,
                    "date_label": row.date_label,
                    "start_date": row.start_date,
                    "end_date": row.end_date,
                }
                for row in source_rows
            ],
        })
    return conflicts
