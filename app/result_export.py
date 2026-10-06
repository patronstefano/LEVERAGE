"""Read-only exports of the saved apparatus classification used by the editor."""

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.display_names import athlete_display_name


HEADERS = [
    'Position', 'Result ID', 'Athlete ID', 'Athlete', 'Country', 'Category',
    'Discipline', 'Format', 'Round', 'Apparatus', 'Day', 'VT attempt',
    'Final Score', 'D Score', 'E Score', 'E est.', 'Penalty', 'Bonus',
    'Event ID', 'Event', 'Year', 'Event start', 'Event end',
]


def export_classification(event, results, file_format):
    rows = []
    for position, result in enumerate(results, 1):
        estimate = None
        if result.E_score is None and result.score is not None and result.D_score is not None:
            value = round(result.score - result.D_score, 3)
            if 0 <= value <= 10:
                estimate = value
        rows.append([
            position, result.id, result.athlete_id, athlete_display_name(result.athlete),
            result.represented_country, result.category.value, result.discipline.value,
            result.format.value, result.round.value, result.apparatus, result.day, result.vt_attempt,
            result.score, result.D_score, result.E_score, estimate, result.Penalty, result.Bonus,
            event.id, event.name, event.year,
            event.start_date.isoformat() if event.start_date else None,
            event.end_date.isoformat() if event.end_date else None,
        ])
    if file_format == 'csv':
        output = StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(HEADERS)
        for row in rows:
            values = []
            for index, value in enumerate(row):
                if value is None:
                    value = ''
                elif 12 <= index <= 17:
                    value = f'{value:.1f}' if index == 13 else f'{value:.3f}'
                elif isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
                    value = "'" + value
                values.append(value)
            writer.writerow(values)
        return output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Classification'
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for cells in sheet.iter_rows():
        for cell in cells:
            # Keep names and imported source text literal, never spreadsheet formulas.
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='191747')
    for cells in sheet.iter_rows(min_row=2, min_col=13, max_col=18):
        for cell in cells:
            cell.number_format = '0.0' if cell.column == 14 else '0.000'
    for index, title in enumerate(HEADERS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 38 if title in {'Athlete', 'Event'} else max(14, len(title) + 2)
    sheet.freeze_panes = 'E2'
    sheet.auto_filter.ref = sheet.dimensions
    notes = workbook.create_sheet('Notes')
    for note in [
        'Saved scores only. Position follows editor ordering, not an official tie-break ranking.',
        'Blank values are unavailable, not zero. E est. = Final Score - D Score; unknown Penalty/Bonus may affect it.',
        'Recorded E is shown separately; it does not certify complete execution data for the whole classification.',
        'Dates describe the event interval, not the exact date of the round.',
    ]:
        notes.append([note])
    notes.column_dimensions['A'].width = 120
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
